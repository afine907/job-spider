"""
数据导出模块

提供 CSV 和 Excel 格式的数据导出功能。
"""

import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Optional, Sequence

from job_spider.storage.models import JobProcessed, JobResponse


class Exporter:
    """
    数据导出器

    支持将职位数据导出为 CSV 或 Excel 格式。
    """

    # 默认导出字段
    DEFAULT_FIELDS = [
        "job_id",
        "source",
        "title",
        "company",
        "salary_min",
        "salary_max",
        "salary_avg",
        "city",
        "district",
        "address",
        "experience",
        "education",
        "skills",
        "company_size",
        "company_industry",
        "company_stage",
        "url",
        "publish_time",
        "crawl_time",
        "status",
    ]

    # 字段中文映射
    FIELD_LABELS = {
        "job_id": "职位ID",
        "source": "数据来源",
        "title": "职位名称",
        "company": "公司名称",
        "salary_min": "最低薪资(K)",
        "salary_max": "最高薪资(K)",
        "salary_avg": "平均薪资(K)",
        "city": "城市",
        "district": "区县",
        "address": "地址",
        "experience": "经验要求",
        "education": "学历要求",
        "description": "职位描述",
        "skills": "技能要求",
        "company_size": "公司规模",
        "company_industry": "行业",
        "company_stage": "融资阶段",
        "url": "链接",
        "publish_time": "发布时间",
        "crawl_time": "爬取时间",
        "status": "状态",
    }

    def __init__(
        self,
        fields: Optional[list[str]] = None,
        include_header: bool = True,
        use_chinese_header: bool = True,
    ) -> None:
        """
        初始化导出器。

        Args:
            fields: 要导出的字段列表（默认使用 DEFAULT_FIELDS）
            include_header: 是否包含表头
            use_chinese_header: 是否使用中文表头
        """
        self.fields = fields or self.DEFAULT_FIELDS.copy()
        self.include_header = include_header
        self.use_chinese_header = use_chinese_header

    def _format_value(self, value: Any) -> str:
        """
        格式化值为字符串。

        Args:
            value: 原始值

        Returns:
            str: 格式化后的字符串
        """
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, Decimal):
            return str(float(value))
        if isinstance(value, list):
            return ", ".join(str(v) for v in value)
        return str(value)

    def _get_header(self) -> list[str]:
        """
        获取表头。

        Returns:
            list[str]: 表头列表
        """
        if self.use_chinese_header:
            return [self.FIELD_LABELS.get(f, f) for f in self.fields]
        return self.fields

    def _extract_row(self, job: JobProcessed | JobResponse | dict) -> list[str]:
        """
        从职位数据中提取一行数据。

        Args:
            job: 职位数据（模型实例或字典）

        Returns:
            list[str]: 数据行
        """
        row = []
        for field in self.fields:
            if isinstance(job, dict):
                value = job.get(field)
            else:
                value = getattr(job, field, None)
            row.append(self._format_value(value))
        return row

    def to_csv(
        self,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        output_path: Optional[str | Path] = None,
        encoding: str = "utf-8-sig",
    ) -> str:
        """
        导出为 CSV 格式。

        Args:
            jobs: 职位数据列表
            output_path: 输出文件路径（可选，不指定则返回字符串）
            encoding: 文件编码

        Returns:
            str: CSV 内容（当 output_path 为 None 时）
        """
        output = StringIO()

        writer = csv.writer(output, quoting=csv.QUOTE_ALL)

        # 写入表头
        if self.include_header:
            writer.writerow(self._get_header())

        # 写入数据
        for job in jobs:
            writer.writerow(self._extract_row(job))

        content = output.getvalue()
        output.close()

        # 写入文件
        if output_path:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding=encoding)
            return str(output_path)

        return content

    def to_csv_bytes(
        self,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        encoding: str = "utf-8-sig",
    ) -> bytes:
        """
        导出为 CSV 格式（字节）。

        Args:
            jobs: 职位数据列表
            encoding: 文件编码

        Returns:
            bytes: CSV 字节数据
        """
        content = self.to_csv(jobs, output_path=None, encoding=encoding)
        return content.encode(encoding)

    def to_excel(
        self,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        output_path: Optional[str | Path] = None,
        sheet_name: str = "职位数据",
    ) -> bytes:
        """
        导出为 Excel 格式。

        需要安装 openpyxl: pip install openpyxl

        Args:
            jobs: 职位数据列表
            output_path: 输出文件路径（可选，不指定则返回字节）
            sheet_name: 工作表名称

        Returns:
            bytes: Excel 字节数据（当 output_path 为 None 时）

        Raises:
            ImportError: 当未安装 openpyxl 时
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError as e:
            raise ImportError(
                "导出 Excel 需要安装 openpyxl: pip install openpyxl"
            ) from e

        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 样式定义
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_alignment = Alignment(vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # 写入表头
        if self.include_header:
            headers = self._get_header()
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

        # 写入数据
        start_row = 2 if self.include_header else 1
        for row_idx, job in enumerate(jobs, start_row):
            row_data = self._extract_row(job)
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = cell_alignment
                cell.border = thin_border

        # 自动调整列宽
        for col_idx, field in enumerate(self.fields, 1):
            header = self._get_header()[col_idx - 1] if self.include_header else field
            # 计算最大宽度
            max_length = len(str(header))
            for row in ws.iter_rows(
                min_row=2 if self.include_header else 1,
                min_col=col_idx,
                max_col=col_idx,
            ):
                for cell in row:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
            # 设置列宽（有上限）
            column_letter = get_column_letter(col_idx)
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        # 冻结首行
        if self.include_header:
            ws.freeze_panes = "A2"

        # 输出
        output = BytesIO()
        wb.save(output)
        content = output.getvalue()
        output.close()

        # 写入文件
        if output_path:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(content)
            return str(output_path)

        return content

    def export(
        self,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        output_path: str | Path,
        format: str = "csv",
        **kwargs,
    ) -> str:
        """
        导出数据到文件。

        Args:
            jobs: 职位数据列表
            output_path: 输出文件路径
            format: 导出格式（csv 或 excel）
            **kwargs: 其他导出参数

        Returns:
            str: 输出文件路径

        Raises:
            ValueError: 不支持的导出格式
        """
        output_path = Path(output_path)

        if format.lower() == "csv":
            return self.to_csv(jobs, output_path=output_path, **kwargs)
        elif format.lower() in ("excel", "xlsx"):
            return self.to_excel(jobs, output_path=output_path, **kwargs)
        else:
            raise ValueError(f"不支持的导出格式: {format}")

    @classmethod
    def quick_export_csv(
        cls,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        output_path: str | Path,
        fields: Optional[list[str]] = None,
    ) -> str:
        """
        快速导出 CSV 的便捷方法。

        Args:
            jobs: 职位数据列表
            output_path: 输出文件路径
            fields: 要导出的字段列表

        Returns:
            str: 输出文件路径
        """
        exporter = cls(fields=fields)
        return exporter.to_csv(jobs, output_path=output_path)

    @classmethod
    def quick_export_excel(
        cls,
        jobs: Sequence[JobProcessed | JobResponse | dict],
        output_path: str | Path,
        fields: Optional[list[str]] = None,
    ) -> str:
        """
        快速导出 Excel 的便捷方法。

        Args:
            jobs: 职位数据列表
            output_path: 输出文件路径
            fields: 要导出的字段列表

        Returns:
            str: 输出文件路径
        """
        exporter = cls(fields=fields)
        return exporter.to_excel(jobs, output_path=output_path)
